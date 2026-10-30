"""Build the monthly management report as Word (.docx) and PDF.

Needs the boardplot CLI, python-docx and, for the PDF, LibreOffice (soffice).
Boardplot draws each exhibit for print: target "print", 160 mm wide (the text
column of A4 with 25 mm margins), PNG at scale 3.125 for 300 dpi. The PNG
records its dpi, so Word places it at 160 mm. Word holds the text, so the
report stays editable. Text numbers come from the specs and the chart reading.
"""
import csv, json, subprocess, tempfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).parent
THEME = json.loads((HERE.parent / "brand-theme.json").read_text())
TMP = Path(tempfile.mkdtemp())
NAVY, BLUE, GREY = RGBColor(0x05, 0x1C, 0x2C), RGBColor(0x22, 0x51, 0xFF), RGBColor(0x5F, 0x6B, 0x7A)


def spec(name):
    return json.loads((HERE / "specs" / f"{name}.json").read_text())


def render(name, s, rows=None, columns=None, height=None):
    """Draw one exhibit at the text width and return (png path, reading)."""
    s["options"]["theme"] = THEME
    path, png = TMP / f"{name}.json", TMP / f"{name}.png"
    path.write_text(json.dumps(s))
    cmd = ["boardplot", "render", str(path), "--target", "print", "--width", "160mm", "--format", "png", "--scale", "3.125",
           "--out", str(png), "--overwrite", "--quiet", "--return", "reading"]
    if height:
        cmd += ["--height", height]
    if rows:
        cmd += ["--rows", str(HERE / "data" / rows), "--columns", columns]
    result = json.loads(subprocess.run(cmd, capture_output=True, text=True).stdout)
    if not result["ok"] or result.get("warnings"):
        raise SystemExit(f"{name}: " + json.dumps(result.get("errors") or result["warnings"], indent=2))
    return png, result["reading"]


def run(paragraph, text, size=10.5, bold=False, color=NAVY):
    r = paragraph.add_run(text)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(size), bold, color
    return r


def para(text="", size=10.5, bold=False, color=NAVY, before=0, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    if text:
        run(p, text, size, bold, color)
    return p


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:fill"), fill)
    pr.append(shd)


def exhibit(number, section, png, text):
    para(section.upper(), 9, True, BLUE, before=6, after=2)
    for line in text:
        para(line)
    para(f"EXHIBIT {number}", 8, True, GREY, before=6, after=2)
    doc.add_picture(str(png))


def page_number(paragraph):
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = paragraph.add_run()
        r.font.name, r.font.size, r.font.color.rgb = "Arial", Pt(8), GREY
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        r._r.append(el)


# Exhibits and the numbers the text needs.
pl = spec("pl-table")
cells = {r["id"]: r.get("cells") for r in pl["data"]["rows"] for r in [r, *r.get("children", [])]}
ebit, marketing, online = cells["ebit"], cells["marketing"], cells["online"]
revenue = {k: sum(cells[c][k] for c in ("online", "wholesale", "stores")) for k in ("mac", "mbu")}
pl_png, _ = render("pl-table", pl)

outlook = spec("revenue-outlook")
values = outlook["data"]["series"][0]["values"]
ac = [v for v in values if isinstance(v, (int, float))]
fc = [v["number"] for v in values if isinstance(v, dict)]
full_year = sum(ac) + sum(fc)
outlook["options"]["title"] = f"The Q4 forecast takes 2026 revenue to €{full_year:.1f}m; December alone brings €{fc[-1]:.1f}m"
outlook_png, _ = render("revenue-outlook", outlook, height="85mm")

pvm = spec("pvm")
pvm_png, reading = render("pvm", pvm, height="85mm")
effects = {d["claimId"]: d["value"] / 1000 for d in reading["derived"]}
pvm["options"]["title"] = (f"Q3 revenue grew €{effects['pvm:delta']:.1f}m: volume added €{effects['pvm:effect:volume']:.1f}m "
                           f"and price €{effects['pvm:effect:price']:.1f}m")
pvm_png, _ = render("pvm", pvm, height="85mm")

bridge_png, bridge = render("ebit-bridge", spec("ebit-bridge"), "ebit-bridge.csv", "label=step,value=eur_k,type=kind", height="85mm")
steps = {r["step"]: float(r["eur_k"]) for r in csv.DictReader((HERE / "data" / "ebit-bridge.csv").read_text().splitlines()) if r["eur_k"]}
ebit_actual = next(d["value"] for d in bridge["derived"] if d["claimId"].startswith("wf:total"))
assert ebit_actual / 1000 == ebit["mac"], "the bridge and the P&L disagree"

# The document.
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.left_margin = section.right_margin = Mm(25)
section.top_margin, section.bottom_margin = Mm(20), Mm(20)
footer = section.footer.paragraphs[0]
run(footer, "Alder Outdoor  |  Monthly management report, September 2026  |  Page ", 8, color=GREY)
page_number(footer)

para("ALDER OUTDOOR", 9, True, BLUE, after=2)
para("Monthly management report", 26, True, after=0)
para("September 2026", 14, color=GREY, after=12)

box = doc.add_table(rows=1, cols=1).cell(0, 0)
shade(box, "F0F3F7")
first = box.paragraphs[0]
run(first, "Executive summary", 12, True)
messages = [
    f"EBIT was €{ebit['mac']:.1f}m, €{ebit['mbu'] - ebit['mac']:.1f}m below budget and €{ebit['mac'] - ebit['mpy']:.1f}m above last year.",
    f"Marketing spent €{marketing['mac'] - marketing['mbu']:.1f}m more than budget on the autumn campaign. This is the whole budget miss.",
    f"Online revenue grew {online['mac'] / online['mpy'] - 1:.0%} on last year to €{online['mac']:.1f}m and carries the group's growth.",
    f"The Q4 forecast takes 2026 revenue to €{full_year:.1f}m. The range for December alone is €{values[-1]['range']['low']:.1f}m to €{values[-1]['range']['high']:.1f}m.",
]
for i, message in enumerate(messages, 1):
    p = box.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(0)
    run(p, f"{i}   ", 10.5, True, BLUE), run(p, message)
box.add_paragraph().paragraph_format.space_after = Pt(0)
exhibit(1, "Results", pl_png, [
    f"September revenue was €{revenue['mac']:.1f}m, {revenue['mac'] / revenue['mbu'] - 1:.0%} above budget. "
    f"Gross margin met budget. Marketing ran {marketing['mac'] / marketing['mbu'] - 1:.0%} over budget, so EBIT missed by "
    f"{1 - ebit['mac'] / ebit['mbu']:.0%}. Year to date, EBIT is {ebit['ytdac'] / ebit['ytdpy'] - 1:.0%} above last year."])
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
exhibit(2, "Revenue outlook", outlook_png, [
    "The forecast for October to December holds the run rate of the last three months and adds the usual December peak. "
    "The bars show the range of the forecast: the wider the bar, the less certain the month."])
exhibit(3, "Revenue drivers", pvm_png, [
    "The price-volume-mix bridge splits the Q3 growth by its causes. Boardplot computes the effects from units and revenue per category."])
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
exhibit(4, "Profit", bridge_png, [
    f"Revenue above budget added €{steps['Revenue above budget']:.0f}k, and a lower margin rate took the same amount back. "
    f"The marketing overspend of €{-steps['Marketing over budget']:.0f}k explains the whole gap to budget."])
para("Next steps", 12, True, before=12)
for action in ["Cap the autumn campaign at its budget for October and November. Owner: CMO.",
               "Re-run the December forecast with the first two weeks of pre-orders. Owner: FP&A.",
               "Report the online growth by country in the October report. Owner: Head of e-commerce."]:
    para(f"–  {action}", after=3)
para("All data in this report is invented. Charts: Boardplot, specs in specs/.", 8, color=GREY, before=12)

props = doc.core_properties
props.author = props.last_modified_by = "Boardplot examples"
props.title = "Monthly management report, September 2026"
doc.save(HERE / "report.docx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(HERE / "report.docx")], check=True, capture_output=True)
print("wrote report.docx and report.pdf")
