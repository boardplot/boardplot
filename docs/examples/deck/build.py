"""Build the example decks. Needs the boardplot CLI and python-pptx.

quarterly-review.pptx is a board deck in a consulting layout. PowerPoint holds
the slide text (tracker, action title, takeaways), so it stays editable.
Boardplot draws each chart for its slot: target "slide", the slot size in pt,
PNG at the default scale 2. The PNG records 192 dpi, so python-pptx places it
at the slot size without scaling.

board-update-after.pptx is board-update-before.pptx with its native charts
replaced by Boardplot charts. The data is read from each old chart, and the
new picture takes the old chart's box.
"""
import csv, json, subprocess, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = Path(__file__).parent
THEME = json.loads((HERE.parent / "brand-theme.json").read_text())
TMP = Path(tempfile.mkdtemp())  # the PNGs end up inside the decks

NAVY, BLUE, GREY, LIGHT, WHITE = (RGBColor.from_string(c) for c in ("051C2C", "2251FF", "5F6B7A", "F0F3F7", "FFFFFF"))
FONT = "Arial"


def render(name, width_pt, height_pt=None, rows=None, columns=None, spec=None):
    """Render one chart for a slide slot and return (png path, reading)."""
    spec = spec or json.loads((HERE / "specs" / f"{name}.json").read_text())
    spec.setdefault("options", {})["theme"] = THEME
    spec_path, png = TMP / f"{name}.json", TMP / f"{name}.png"
    spec_path.write_text(json.dumps(spec))
    cmd = ["boardplot", "render", str(spec_path), "--target", "slide", "--format", "png",
           "--width", f"{width_pt}pt", "--out", str(png), "--overwrite", "--quiet", "--return", "reading"]
    if height_pt:
        cmd += ["--height", f"{height_pt}pt"]
    if rows:
        cmd += ["--rows", str(rows), "--columns", columns]
    result = json.loads(subprocess.run(cmd, capture_output=True, text=True).stdout)
    if not result["ok"] or result.get("warnings"):
        raise SystemExit(f"{name}: " + json.dumps(result.get("errors") or result["warnings"], indent=2))
    return png, result["reading"]


def text(slide, s, x, y, w, h, size=11, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = box.text_frame
    frame.word_wrap, frame.vertical_anchor = True, anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(s.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(size), bold, color
    return box


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    r = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    style = r._element.find(qn("p:style"))  # no theme shadow or outline
    r._element.remove(style)
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.fill.background()
    return r


def content_slide(prs, tracker, title, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, tracker.upper(), 36, 22, 600, 12, size=9, bold=True, color=BLUE)
    text(slide, title, 36, 38, 888, 56, size=20, bold=True)
    rect(slide, 36, 98, 888, 0.75, NAVY)
    text(slide, "Alder Outdoor  |  Q3 FY26 business review", 36, 516, 400, 12, size=8, color=GREY)
    text(slide, str(page), 888, 516, 36, 12, size=8, color=GREY, align=PP_ALIGN.RIGHT)
    return slide


def takeaways(slide, points, x=636, y=112, w=288, h=376):
    rect(slide, x, y, w, h, LIGHT)
    text(slide, "What it means", x + 16, y + 16, w - 32, 16, size=12, bold=True)
    top = y + 48
    for n, point in enumerate(points, start=1):
        circle = rect(slide, x + 16, top, 18, 18, NAVY, MSO_SHAPE.OVAL)
        circle.text_frame.text = str(n)
        run = circle.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(9), True, WHITE
        circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        text(slide, point, x + 44, top, w - 60, 70, size=11)
        top += 80


def picture(slide, png, x, y):
    slide.shapes.add_picture(str(png), Pt(x), Pt(y))  # placed at its recorded dpi


def save(prs, name, title):
    props = prs.core_properties
    props.author = props.last_modified_by = "Boardplot examples"
    props.title = title
    prs.save(HERE / name)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)  # 16:9
    return prs


def quarterly_review():
    prs = new_deck()
    data = HERE / "data"

    # 1 Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 960, 540, NAVY)
    rect(slide, 36, 200, 48, 4, BLUE)
    text(slide, "Q3 FY26 business review", 36, 216, 700, 48, size=36, bold=True, color=WHITE)
    text(slide, "Board of directors  |  14 October 2026", 36, 272, 700, 20, size=14, color=WHITE)
    text(slide, "Alder Outdoor. All figures are invented for the Boardplot examples.", 36, 500, 700, 14, size=9, color=WHITE)

    # 2 Executive summary: four KPI tiles and three messages
    slide = content_slide(prs, "Executive summary", "Revenue is 9% above last year and ahead of plan; margin and costs need attention", 2)
    for i, name in enumerate(["kpi-revenue", "kpi-margin", "kpi-ebit", "kpi-online"]):
        png, _ = render(name, 210)
        picture(slide, png, 36 + i * 226, 112)
    messages = [
        ("Growth comes from volume, price and the new bike packs", "Online grew 27%. Dropping climbing hardware cost €1.1m of revenue."),
        ("Gross margin fell 1.3 points to 51.2%", "Freight, duties and a richer marketing push grew faster than revenue."),
        ("Two transformation workstreams are off track", "Sourcing in Asia and the central warehouse need a decision before the roll-out."),
    ]
    for i, (head, body) in enumerate(messages):
        top = 318 + i * 62
        rect(slide, 36, top, 4, 48, BLUE)
        text(slide, head, 52, top, 860, 18, size=13, bold=True)
        text(slide, body, 52, top + 22, 860, 18, size=11, color=GREY)

    # 3 Revenue bridge, price-volume-mix
    png, reading = render("pvm", 576, 376)
    # The title takes its numbers from the effects Boardplot computed.
    effects = {d["claimId"]: d["value"] for d in reading["derived"]}
    growth = effects["pvm:delta"]
    slide = content_slide(prs, "Revenue", f"Q3 revenue grew €{growth / 1000:.1f}m: volume added €{effects['pvm:effect:volume'] / 1000:.1f}m and price €{effects['pvm:effect:price'] / 1000:.1f}m", 3)
    picture(slide, png, 36, 112)
    takeaways(slide, [
        "More units of tents, backpacks, footwear and accessories carry the growth. Apparel lost volume to discounters.",
        "Price increases on tents and footwear held without a volume loss.",
        "Bike packs, launched in May, bring €1.5m a quarter and replace the dropped climbing range.",
    ])

    # 4 Channel shift with CAGR per channel
    png, _ = render("channel-mix", 576, 376, data / "channel-revenue.csv", "category=year,series=channel,value=revenue_eurm")
    slide = content_slide(prs, "Channels", "Online grew 28.5% a year since FY21 and brought four fifths of the group's growth", 4)
    picture(slide, png, 36, 112)
    takeaways(slide, [
        "Online nearly tripled from €22m to €60m in four years.",
        "Own stores have been flat since FY22. Each store needs a clear role next to online.",
        "Wholesale grows slowly but stays profitable. Keep the key accounts.",
    ])

    # 5 Regions: two charts side by side
    slide = content_slide(prs, "Regions", "DACH and Iberia lead the regional growth; France fell 8% and is the only region below last year", 5)
    rows, columns = data / "region-revenue.csv", "category=region,series=scenario,value=revenue_eurm"
    png, _ = render("region-variance", 436, 376, rows, columns)
    picture(slide, png, 36, 112)
    png, _ = render("region-growth", 436, 376, rows, columns)
    picture(slide, png, 488, 112)

    # 6 Market position
    png, _ = render("market-map", 576, 376)
    slide = content_slide(prs, "Market", "Alder leads the Nordics with 27% share but holds less than 10% in France and the UK", 6)
    picture(slide, png, 36, 112)
    takeaways(slide, [
        "In the Nordics, 27% share gives Alder pricing power. Protect it.",
        "France and the UK are large markets where Alder has 8%. Growth there needs online first.",
        "Summit Co. leads in the three largest markets.",
    ])

    # 7 P&L
    png, _ = render("pl-table", 888)
    slide = content_slide(prs, "Profit and loss", "EBIT rose 11% to €6.8m but missed plan by 8%: product cost, freight and marketing grew faster than revenue", 7)
    picture(slide, png, 36, 112)

    # 8 Transformation programme
    png, _ = render("programme", 576, 376)
    slide = content_slide(prs, "Transformation", "Sourcing in Asia and the central warehouse are off track in their pilots", 8)
    picture(slide, png, 36, 112)
    takeaways(slide, [
        "Sourcing: the second supplier failed the quality audit. Decide on a third supplier by November.",
        "Warehouse: the pilot runs 6 weeks late. The roll-out moves to Q2 FY27.",
        "Pricing: the roll-out is at risk while store staff wait for training.",
    ])

    save(prs, "quarterly-review.pptx", "Q3 FY26 business review")


def board_update_before():
    """The deck to revise: a typical slide with two default PowerPoint charts."""
    regions = list(csv.DictReader((HERE / "data" / "region-revenue.csv").read_text().splitlines()))
    names = list(dict.fromkeys(r["region"] for r in regions))
    value = {(r["region"], r["scenario"]): float(r["revenue_eurm"]) for r in regions}

    prs = new_deck()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Regional performance Q3"
    chart_data = CategoryChartData()
    chart_data.categories = names
    chart_data.add_series("Q3 FY26", [value[(n, "AC")] for n in names])
    chart_data.add_series("Q3 FY25", [value[(n, "PY")] for n in names])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(36), Pt(130), Pt(540), Pt(360), chart_data).chart
    chart.has_legend, chart.legend.position, chart.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False

    chart_data = CategoryChartData()
    chart_data.categories = ["Own stores", "Wholesale", "Online"]
    chart_data.add_series("Revenue Q3", [24.2, 15.1, 20.3])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Pt(600), Pt(130), Pt(324), Pt(360), chart_data).chart
    chart.has_legend, chart.legend.position = True, XL_LEGEND_POSITION.BOTTOM
    # Name each chart's embedded workbook after what it holds.
    for n, shape in enumerate([s for s in slide.shapes if s.has_chart], start=1):
        shape.chart.part.chart_workbook.xlsx_part.partname = PackURI(f"/ppt/embeddings/chart-data-{n}.xlsx")
    save(prs, "board-update-before.pptx", "Board update")


def board_update_after():
    """Replace each native chart with a Boardplot chart in the same box."""
    prs = Presentation(HERE / "board-update-before.pptx")
    for slide in prs.slides:
        for n, shape in enumerate([s for s in slide.shapes if s.has_chart], start=1):
            plot = shape.chart.plots[0]
            # Read the data from the old chart.
            rows = [{"category": c, "series": s.name, "value": v}
                    for s in plot.series for c, v in zip(plot.categories, s.values)]
            rows_file = TMP / f"rows-{n}.json"
            rows_file.write_text(json.dumps(rows))
            spec = json.loads((HERE / "specs" / f"board-update-{n}.json").read_text())
            png, _ = render(f"board-update-{n}", round(shape.width.pt), round(shape.height.pt),
                            rows_file, "category=category,series=series,value=value", spec)
            picture(slide, png, shape.left.pt, shape.top.pt)
            # Remove the old chart with its chart part and workbook.
            slide.part.drop_rel(shape._element.chart_rId)
            shape._element.getparent().remove(shape._element)
    save(prs, "board-update-after.pptx", "Board update")


if __name__ == "__main__":
    quarterly_review()
    board_update_before()
    board_update_after()
